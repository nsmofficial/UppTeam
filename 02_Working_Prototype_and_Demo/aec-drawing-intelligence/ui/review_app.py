"""Engineer review screen.

    streamlit run ui/review_app.py -- out/seed_00      (or: aecdi review out/seed_00)

Findings sorted by tier and severity; click one to see every evidence anchor cropped from the
source sheet / spec page / BOQ row; Accept / Reject (with a reason) / Comment. Decisions are
appended to review_log.jsonl (who, when, what, why) - the audit trail. The tool never edits
drawings and never approves anything.
"""
from __future__ import annotations

import datetime as dt
import json
import sys
from pathlib import Path

import pandas as pd
import pymupdf
import streamlit as st
from openpyxl import load_workbook

OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[1] / "out"
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from aecdi.models import Finding  # noqa: E402  (stable fingerprints for the audit log)

DISCIPLINE_NAMES = {"G": "General", "A": "Architectural", "S": "Structural", "M": "Mechanical", "E": "Electrical",
                    "P": "Plumbing", "FP": "Fire protection", "BOQ": "BOQ"}
REJECT_REASONS = ["Not a discrepancy (design intent)", "Duplicate of another finding", "Evidence points to wrong place",
                  "Already resolved / superseded", "Other"]
TIER_ICON = {"confirmed": "🔴 confirmed", "probable": "🟠 probable", "needs_review": "⚪ needs review"}

st.set_page_config(page_title="AEC Drawing Intelligence - Review", layout="wide")


@st.cache_data
def load_result(path: str) -> dict:
    return json.loads((Path(path) / "findings.json").read_text(encoding="utf-8"))


def load_log() -> list[dict]:
    p = OUT / "review_log.jsonl"
    if not p.exists():
        return []
    return [json.loads(l) for l in p.read_text(encoding="utf-8").splitlines() if l.strip()]


def append_log(entry: dict):
    with (OUT / "review_log.jsonl").open("a", encoding="utf-8") as fh:
        fh.write(json.dumps(entry) + "\n")


@st.cache_data
def crop(doc: str, page: int, bbox: tuple, pad: float = 70) -> bytes:
    d = pymupdf.open(OUT / "sources" / doc)
    pg = d[page]
    r = pymupdf.Rect(bbox)
    clip = pymupdf.Rect(r.x0 - pad * 2, r.y0 - pad, r.x1 + pad * 2, r.y1 + pad) & pg.rect
    pg.draw_rect(r + (-3, -3, 3, 3), color=(0.9, 0.1, 0.1), width=1.6)
    return pg.get_pixmap(dpi=150, clip=clip).tobytes("png")


@st.cache_data
def boq_row(doc: str, sheet: str, row: int) -> pd.DataFrame:
    ws = load_workbook(OUT / "sources" / doc, data_only=False)[sheet]
    hdr_row = next(r for r in range(1, 30) if any(str(c.value or "").upper().startswith("ITEM") for c in ws[r]))
    headers = [str(c.value or "") for c in ws[hdr_row]]
    vals = [c.value for c in ws[row]]
    return pd.DataFrame([vals], columns=headers, index=[f"row {row}"])


res = load_result(str(OUT))
findings = res["findings"]
for f in findings:
    f["fingerprint"] = Finding(**{k: v for k, v in f.items() if k != "fingerprint"}).fingerprint()
    f["discipline_names"] = sorted({DISCIPLINE_NAMES.get(d, d) for d in f["disciplines"]}
                                   | ({"BOQ"} if "BOQ" in f["sheets"] else set())) or ["General"]
log = load_log()
by_fp = {}
for e in log:
    if e.get("decision") in ("accept", "reject"):
        by_fp[e.get("fingerprint") or e["finding_id"]] = e
latest = {f["id"]: by_fp[f["fingerprint"]] for f in findings if f["fingerprint"] in by_fp}
latest.update({f["id"]: by_fp[f["id"]] for f in findings if f["id"] in by_fp and f["id"] not in latest})

# ---------------------------------------------------------------- sidebar
st.sidebar.title("AEC Drawing Intelligence")
st.sidebar.caption(f"Project: **{res['project']}**  \nOutput: `{OUT}`")
reviewer = st.sidebar.text_input("Reviewer", value=st.session_state.get("reviewer", "QC reviewer"))
st.session_state["reviewer"] = reviewer
tiers = st.sidebar.multiselect("Tier", ["confirmed", "probable", "needs_review"], ["confirmed", "probable", "needs_review"])
sevs = st.sidebar.multiselect("Severity", ["high", "medium", "low"], ["high", "medium", "low"])
types = sorted({f["type"] for f in findings})
sel_types = st.sidebar.multiselect("Type", types, types)
discs = sorted({d for f in findings for d in f["discipline_names"]})
sel_discs = st.sidebar.multiselect("Discipline", discs, discs)
status_filter = st.sidebar.radio("Status", ["All", "Open", "Accepted", "Rejected"], horizontal=True)
st.sidebar.divider()
st.sidebar.caption("Findings are *potential discrepancies*. Every evidence anchor shown has been re-verified "
                   "against the source by the evidence gate. The engineer of record decides.")

s = res["stats"]
c = st.columns(6)
c[0].metric("Findings", s["findings"])
c[1].metric("Needs review", s["needs_review"])
c[2].metric("Rejected by gate", s["rejected_by_gate"])
c[3].metric("Anchors verified", f"{s['anchors_verified']}/{s['anchors_checked']}")
c[4].metric("Reviewed", f"{len(latest)}/{len(findings)}")
c[5].metric("Run time", f"{s['seconds']} s")

tab_f, tab_rev, tab_mk, tab_gate, tab_acc, tab_eval, tab_exp = st.tabs(
    ["Findings", "Revision delta", "Redline closure", "Blocked by evidence gate", "ACC issues", "Evaluation", "Exports"])

with tab_f:
    def status_of(fid):
        return {"accept": "Accepted", "reject": "Rejected"}.get(latest.get(fid, {}).get("decision"), "Open")

    rows = [f for f in findings if f["tier"] in tiers and f["severity"] in sevs and f["type"] in sel_types
            and set(f["discipline_names"]) & set(sel_discs)
            and (status_filter == "All" or status_of(f["id"]) == status_filter)]
    df = pd.DataFrame([{"ID": f["id"], "Tier": TIER_ICON[f["tier"]], "Sev": f["severity"], "Type": f["type"].replace("_", " "),
                        "Sheets": ", ".join(f["sheets"]), "Potential discrepancy": f["statement"],
                        "Status": status_of(f["id"])} for f in rows])
    left, right = st.columns([5, 6])
    with left:
        if df.empty:
            st.info("No findings match the filters.")
            sel = None
        else:
            ev = st.dataframe(df, hide_index=True, use_container_width=True, height=640, on_select="rerun",
                              selection_mode="single-row", key="findings_table",
                              column_config={"Potential discrepancy": st.column_config.TextColumn(width="large")})
            if ev.selection.rows:
                st.session_state["sel_id"] = rows[ev.selection.rows[0]]["id"]
            sel_id = st.session_state.get("sel_id")
            sel = next((f for f in rows if f["id"] == sel_id), rows[0])
    with right:
        if sel:
            f = sel
            st.subheader(f"{f['id']} · {f['type'].replace('_', ' ')}")
            st.markdown(f"**{TIER_ICON[f['tier']]}** · severity **{f['severity']}** · `{f['method']}` · "
                        f"confidence {f['confidence']:.2f}")
            st.write(f["statement"])
            if f.get("rationale"):
                st.caption(f"Why this tier: {f['rationale']}")
            st.markdown("**Evidence** (each anchor re-found in the source)")
            for e in f["evidence"]:
                label = f"{e['role']} — {e['doc']}"
                if e.get("xlsx_sheet"):
                    st.markdown(f"*{label}* · `{e['xlsx_sheet']}!{e['cell']}` = **{e['snippet']}**")
                    st.dataframe(boq_row(e["doc"], e["xlsx_sheet"], e["row"]), use_container_width=True)
                elif e.get("row") and e.get("cell") and not e.get("xlsx_sheet"):
                    st.markdown(f"*{label}* · line {e['row']}, column `{e['cell']}` = **{e['snippet'] or '(blank)'}**")
                elif e.get("page") is not None and e.get("bbox"):
                    st.markdown(f"*{label}* · {e.get('sheet')} (page {e['page'] + 1}) · “{e['snippet']}”")
                    st.image(crop(e["doc"], e["page"], tuple(e["bbox"])))
            st.divider()
            cur = latest.get(f["id"])
            if cur:
                st.info(f"Last decision: **{cur['decision']}** by {cur['reviewer']} at {cur['at']}"
                        + (f" — {cur.get('reason')}" if cur.get("reason") else "")
                        + (f" — “{cur.get('comment')}”" if cur.get("comment") else ""))
            comment = st.text_area("Comment", key=f"c_{f['id']}")
            reason = st.selectbox("Reject reason", REJECT_REASONS, key=f"r_{f['id']}")
            b1, b2, b3 = st.columns(3)
            ids = [x["id"] for x in rows]
            nxt = next((i for i in ids[ids.index(f["id"]) + 1:] if status_of(i) == "Open"), f["id"])

            def decide(d: dict):
                append_log(d)
                st.session_state["sel_id"] = nxt          # auto-advance to the next open finding
                st.session_state.pop("findings_table", None)
                st.rerun()

            entry = {"finding_id": f["id"], "fingerprint": f["fingerprint"], "sheets": f["sheets"],
                     "type": f["type"], "entities": f["entities"], "method": f["method"],
                     "tier": f["tier"], "reviewer": reviewer, "at": dt.datetime.now().isoformat(timespec="seconds"),
                     "comment": comment}
            if b1.button("✅ Accept", key=f"a_{f['id']}", use_container_width=True):
                decide(entry | {"decision": "accept"})
            if b2.button("❌ Reject", key=f"x_{f['id']}", use_container_width=True):
                decide(entry | {"decision": "reject", "reason": reason})
            if b3.button("💬 Comment only", key=f"m_{f['id']}", use_container_width=True):
                append_log(entry | {"decision": "comment"})
                st.rerun()

with tab_rev:
    ch = res.get("revision_changes", [])
    st.write(f"{len(ch)} change(s) between revisions. Unclouded changes are also listed as findings.")
    if ch:
        st.dataframe(pd.DataFrame([{"Sheet": x["sheet"], "Kind": x["kind"], "Before": x["before"], "After": x["after"],
                                    "Clouded": "yes" if x["clouded"] else "NO"} for x in ch]),
                     hide_index=True, use_container_width=True)
    ov = sorted((OUT / "overlays").glob("*.png")) if (OUT / "overlays").exists() else []
    if ov:
        pick = st.selectbox("Overlay (red = only in previous revision, blue = only in current)", [p.name for p in ov])
        st.image(str(OUT / "overlays" / pick))
        st.caption("A pixel overlay exaggerates change when rows are inserted (everything below shifts). The delta "
                   "table above is the semantic diff: schedules by row key, text and geometry with layout-shift compensation.")

with tab_mk:
    mk = res.get("markups", [])
    if not mk:
        st.info("No reviewer redlines were supplied (add a *_REDLINES.pdf of the previous revision to the project folder).")
    else:
        st.write(f"{len(mk)} reviewer markup(s) on the previous revision; "
                 f"{sum(1 for m in mk if not m['incorporated'])} without a matching change in the current revision.")
        st.dataframe(pd.DataFrame([{"Sheet": m["sheet"], "Markup": m["text"], "Row": m.get("row") or "",
                                    "Column": m.get("column") or "", "Incorporated": "yes" if m["incorporated"] else "NO",
                                    "Matched change": "; ".join(m["matched_changes"])} for m in mk]),
                     hide_index=True, use_container_width=True)

with tab_acc:
    st.write("Issues that would be sent to Autodesk Construction Cloud. Accepted findings are posted as *open*; "
             "LLM and probable findings are never posted unless accepted here.")
    policy = st.radio("Policy", ["accepted", "accepted+confirmed-draft"], horizontal=True)
    from aecdi.integrations.acc import load_review_log, payloads
    items = payloads([Finding(**{k: v for k, v in f.items() if k not in ("fingerprint", "discipline_names")}) for f in findings],
                     load_review_log(OUT), policy=policy)
    st.write(f"{len(items)} issue(s) selected.")
    for it in items:
        with st.expander(f"{it['finding_id']} · {it['body']['status']} · {it['body']['title']}"):
            st.code(json.dumps(it["body"], indent=2), language="json")
    st.caption("Post with: `aecdi acc-post <out> --project <ACC project id> --subtype <issue subtype id> --live` "
               "(needs a 3-legged APS token with data:write in $APS_TOKEN).")

with tab_gate:
    rej = res.get("rejected", [])
    st.write(f"{len(rej)} candidate finding(s) were blocked because their evidence could not be re-found in the source.")
    for f in rej:
        with st.expander(f"{f['id']} · {f['method']} · {f['statement'][:110]}"):
            st.write(f["rationale"])
            st.json([{k: e.get(k) for k in ("doc", "sheet", "page", "snippet", "verified", "role")} for e in f["evidence"]])

with tab_eval:
    rep = ROOT / "eval" / "report.md"
    if not rep.exists():
        rep = ROOT.parents[1] / "03_Evaluation_Plan_and_Hallucination_Detection" / "benchmark_data" / "report.md"
    if rep.exists():
        st.markdown(rep.read_text(encoding="utf-8"))
    else:
        st.info("Run `aecdi eval` to produce eval/report.md")
    if log:
        st.markdown("#### Live precision from reviewer decisions (this project)")
        dec = [e for e in latest.values() if e["decision"] in ("accept", "reject")]
        if dec:
            acc = sum(1 for e in dec if e["decision"] == "accept")
            st.metric("Accept rate (live precision proxy)", f"{acc}/{len(dec)} = {acc / len(dec):.0%}")

with tab_exp:
    for name in ["QC_Issue_Log.xlsx", "CD_Set_RevB_REVIEW.pdf", "Project_Manual_REVIEW.pdf", "acc_issues.json",
                 "findings.json", "project_graph.json", "review_log.jsonl"]:
        p = OUT / name
        if p.exists():
            st.download_button(f"Download {name}", p.read_bytes(), file_name=name, key=f"dl_{name}")
