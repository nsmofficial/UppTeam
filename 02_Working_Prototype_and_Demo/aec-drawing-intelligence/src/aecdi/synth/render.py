"""Render a synthetic Project into vector PDFs (drawing set + project manual), an Excel BOQ and
an answer key. Drawing conventions follow common U.S. practice: 11x17 half-size sheets,
1/8" = 1'-0" plans, right-hand title block, grid bubbles, revision clouds with deltas.
"""
from __future__ import annotations

import json
import math
from pathlib import Path

import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ..units import format_feet_inches
from .project import (MM_PER_FT, fmt_len, DOOR_TYPES, FIRM_NAME, PROJECT_NAME, PROJECT_NO, SHEETS, SPEC_REFRIGERANT, Project,
                      make_project)

W, H = 1224.0, 792.0          # 17 x 11 in
SCALE = 9.0                   # pt per ft at 1/8" = 1'-0"  (set per project by render_set)
SCALE_IMPERIAL = 9.0
SCALE_METRIC = 72 / 25.4 * MM_PER_FT / 100      # 1:100 -> 8.64 pt per ft


def scale_label(p) -> str:
    return "1:100" if p.units == "metric" else "1/8\" = 1'-0\""
OX, OY = 150.0, 170.0         # plan origin = grid 1 / grid A
TB = pymupdf.Rect(1040, 18, 1210, 774)
BLACK, GREY, RED = (0, 0, 0), (0.55, 0.55, 0.55), (0.85, 0.05, 0.05)


class Canvas:
    """Thin wrapper that draws and records the bbox of every element id (for the answer key)."""

    def __init__(self, page: pymupdf.Page):
        self.page = page
        self.elements: dict[tuple, pymupdf.Rect] = {}

    def rec(self, eid, rect):
        if eid is None:
            return
        r = pymupdf.Rect(rect)
        self.elements[eid] = self.elements[eid] | r if eid in self.elements else r

    def text(self, x, y, s, size=7.0, bold=False, eid=None, rotate=0, color=BLACK, align="left"):
        font = "hebo" if bold else "helv"
        w = pymupdf.get_text_length(s, fontname=font, fontsize=size)
        if rotate == 90:
            if align == "center":
                y += w / 2
            self.page.insert_text((x, y), s, fontsize=size, fontname=font, color=color, rotate=90)
            r = pymupdf.Rect(x - size * 0.95, y - w, x + size * 0.25, y)
        else:
            if align == "center":
                x -= w / 2
            elif align == "right":
                x -= w
            self.page.insert_text((x, y), s, fontsize=size, fontname=font, color=color)
            r = pymupdf.Rect(x, y - size * 0.95, x + w, y + size * 0.25)
        self.rec(eid, r)
        return r

    def line(self, p1, p2, width=0.6, color=BLACK, dashes=None, eid=None):
        sh = self.page.new_shape()
        sh.draw_line(p1, p2)
        sh.finish(color=color, width=width, dashes=dashes)
        sh.commit()
        self.rec(eid, pymupdf.Rect(min(p1[0], p2[0]), min(p1[1], p2[1]), max(p1[0], p2[0]), max(p1[1], p2[1])))

    def rect(self, r, width=0.6, color=BLACK, fill=None, dashes=None, eid=None):
        sh = self.page.new_shape()
        sh.draw_rect(r)
        sh.finish(color=color, width=width, fill=fill, dashes=dashes)
        sh.commit()
        self.rec(eid, r)

    def circle(self, c, radius, width=0.6, color=BLACK, eid=None):
        sh = self.page.new_shape()
        sh.draw_circle(c, radius)
        sh.finish(color=color, width=width)
        sh.commit()
        self.rec(eid, pymupdf.Rect(c[0] - radius, c[1] - radius, c[0] + radius, c[1] + radius))

    def arc(self, p1, corner, p3, width=0.4, eid=None):
        sh = self.page.new_shape()
        sh.draw_curve(p1, corner, p3)
        sh.finish(color=BLACK, width=width)
        sh.commit()
        self.rec(eid, pymupdf.Rect(p1, p1) | p3 | corner)

    def cloud(self, r: pymupdf.Rect, rev: str, margin=6.0, scallop=9.0):
        r = pymupdf.Rect(r.x0 - margin, r.y0 - margin, r.x1 + margin, r.y1 + margin)
        corners = [r.tl, r.tr, r.br, r.bl, r.tl]
        sh = self.page.new_shape()
        for a, b in zip(corners, corners[1:]):
            length = math.dist(a, b)
            n = max(2, round(length / scallop))
            dx, dy = (b.x - a.x) / n, (b.y - a.y) / n
            nx, ny = dy / (length / n), -dx / (length / n)        # outward normal (clockwise path)
            for i in range(n):
                p = pymupdf.Point(a.x + dx * i, a.y + dy * i)
                q = pymupdf.Point(a.x + dx * (i + 1), a.y + dy * (i + 1))
                c = pymupdf.Point((p.x + q.x) / 2 + nx * scallop * 0.55, (p.y + q.y) / 2 + ny * scallop * 0.55)
                sh.draw_curve(p, c, q)
        sh.finish(color=RED, width=0.9)
        sh.commit()
        # delta symbol at top-right
        tx, ty = r.x1 + 4, r.y0 - 4
        sh = self.page.new_shape()
        sh.draw_polyline([(tx, ty), (tx + 14, ty), (tx + 7, ty - 12), (tx, ty)])
        sh.finish(color=RED, width=0.8, closePath=True)
        sh.commit()
        self.text(tx + 7, ty - 2, rev, size=7, bold=True, color=RED, align="center")


# ----------------------------------------------------------------------------------------------
def title_block(cv: Canvas, p: Project, sheet: str, title: str, scale: str):
    cv.rect(TB, width=1.2)
    x = TB.x0 + 8
    cv.text(x, 34, FIRM_NAME, size=8.5, bold=True)
    cv.text(x, 46, "SYNTHETIC SAMPLE - NOT FOR CONSTRUCTION", size=5.5, color=RED)
    cv.line((TB.x0, 54), (TB.x1, 54))
    cv.text(x, 68, "PROJECT:", size=6)
    cv.text(x, 80, PROJECT_NAME, size=7.5, bold=True)
    cv.text(x, 94, f"PROJECT NO.: {PROJECT_NO}", size=6)
    cv.line((TB.x0, 104), (TB.x1, 104))
    cv.text(x, 118, "REVISIONS", size=6.5, bold=True)
    cv.text(x, 132, "REV", size=6, bold=True)
    cv.text(x + 24, 132, "DATE", size=6, bold=True)
    cv.text(x + 74, 132, "DESCRIPTION", size=6, bold=True)
    for i, (rev, date, desc) in enumerate(p.revisions[sheet]):
        y = 146 + 13 * i
        cv.text(x + 2, y, rev, size=6.5, eid=("title_block", ""))
        cv.text(x + 24, y, date, size=6, eid=("title_block", ""))
        cv.text(x + 74, y, desc, size=6, eid=("title_block", ""))
    cv.line((TB.x0, 300), (TB.x1, 300))
    cv.text(x, 314, "SHEET TITLE:", size=6)
    words, lines, cur = title.split(), [], ""
    for wd in words:
        if len(cur) + len(wd) + 1 > 26:
            lines.append(cur)
            cur = wd
        else:
            cur = f"{cur} {wd}".strip()
    lines.append(cur)
    for i, ln in enumerate(lines):
        cv.text(x, 328 + 11 * i, ln, size=8, bold=True)
    cv.text(x, 368, "SCALE:", size=6)
    cv.text(x, 380, scale, size=7)
    cv.line((TB.x0, 690), (TB.x1, 690))
    cv.text(x, 704, "SHEET NO.", size=6)
    cv.text(x, 752, sheet, size=26, bold=True, eid=("title_block", ""))


def grid_positions(p: Project):
    xs = [OX]
    for b in p.bays_x:
        xs.append(xs[-1] + b * SCALE)
    ys = [OY]
    for b in p.bays_y:
        ys.append(ys[-1] + b * SCALE)
    return xs, ys


def draw_grids(cv: Canvas, p: Project):
    xs, ys = grid_positions(p)
    bottom = ys[-1] + 20
    right = xs[-1] + 20
    for i, gx in enumerate(xs):
        cv.line((gx, OY - 67), (gx, bottom), width=0.35, color=GREY, dashes="[6 2 1 2] 0")
        cv.circle((gx, OY - 78), 11)
        cv.text(gx, OY - 75, str(i + 1), size=9, bold=True, align="center")
    for j, gy in enumerate(ys):
        cv.line((OX - 67, gy), (right, gy), width=0.35, color=GREY, dashes="[6 2 1 2] 0")
        cv.circle((OX - 78, gy), 11)
        cv.text(OX - 78, gy + 3, "ABCDE"[j], size=9, bold=True, align="center")


def dim_h(cv: Canvas, x0, x1, y, label, eid=None):
    cv.line((x0, y), (x1, y), width=0.4)
    for x in (x0, x1):
        cv.line((x, y - 5), (x, y + 5), width=0.3)
        cv.line((x - 2.5, y + 2.5), (x + 2.5, y - 2.5), width=0.8)
    cv.text((x0 + x1) / 2, y - 3, label, size=7, align="center", eid=eid)


def dim_v(cv: Canvas, x, y0, y1, label, eid=None):
    cv.line((x, y0), (x, y1), width=0.4)
    for y in (y0, y1):
        cv.line((x - 5, y), (x + 5, y), width=0.3)
        cv.line((x - 2.5, y + 2.5), (x + 2.5, y - 2.5), width=0.8)
    cv.text(x - 3, (y0 + y1) / 2, label, size=7, rotate=90, align="center", eid=eid)


def draw_dims(cv: Canvas, p: Project, sheet: str):
    xs, ys = grid_positions(p)
    for i, b in enumerate(p.bays_x):
        label = p.dim_override.get((sheet, i), fmt_len(p.units, b))
        dim_h(cv, xs[i], xs[i + 1], OY - 28, label, eid=("dim", f"x{i}"))
    dim_h(cv, xs[0], xs[-1], OY - 48, fmt_len(p.units, p.width), eid=("dim", "overall_x"))
    for j, b in enumerate(p.bays_y):
        dim_v(cv, OX - 28, ys[j], ys[j + 1], fmt_len(p.units, b), eid=("dim", f"y{j}"))
    dim_v(cv, OX - 48, ys[0], ys[-1], fmt_len(p.units, p.depth), eid=("dim", "overall_y"))


def P(x_ft, y_ft):
    return (OX + x_ft * SCALE, OY + y_ft * SCALE)


def draw_building(cv: Canvas, p: Project, light=False, doors=True):
    col = GREY if light else BLACK
    wt = 0.5 if light else 1.6
    cy, ch = p.corridor_y, p.corridor_half
    ent = 3.0  # half opening of the entry pairs
    x0, y0 = P(0, 0)
    x1, y1 = P(p.width, p.depth)
    cv.line((x0, y0), (x1, y0), width=wt, color=col)
    cv.line((x0, y1), (x1, y1), width=wt, color=col)
    for xe in (x0, x1):
        cv.line((xe, y0), (xe, OY + (cy - ent) * SCALE), width=wt, color=col)
        cv.line((xe, OY + (cy + ent) * SCALE), (xe, y1), width=wt, color=col)
    for side, yw in (("N", cy - ch), ("S", cy + ch)):
        gaps = sorted(d.x for d in p.doors if d.side == side and d.on_plan and d.x is not None) if doors else []
        cur = 0.0
        for gx in gaps:
            cv.line(P(cur, yw), P(gx - 1.5, yw), width=wt * 0.8, color=col)
            cur = gx + 1.5
        cv.line(P(cur, yw), P(p.width, yw), width=wt * 0.8, color=col)
    for r in p.rooms:
        if r.side in ("N", "S") and r.x1 < p.width - 0.01:
            cv.line(P(r.x1, r.y0), P(r.x1, r.y1), width=wt * 0.8, color=col)


def draw_doors(cv: Canvas, p: Project):
    cy, ch = p.corridor_y, p.corridor_half
    for d in p.doors:
        if not d.on_plan or d.x is None:
            continue
        label = d.plan_label or d.mark
        g = ("door_group", d.mark)
        if d.side in ("N", "S"):
            yw = cy - ch if d.side == "N" else cy + ch
            into = -1 if d.side == "N" else 1                 # swing into the room
            hinge = P(d.x - 1.5, yw)
            leaf_end = P(d.x - 1.5, yw + into * 3.0)
            strike = P(d.x + 1.5, yw)
            cv.line(hinge, leaf_end, width=0.8, eid=g)
            cv.arc(leaf_end, P(d.x + 1.5, yw + into * 3.0), strike, eid=g)
            tx, ty = P(d.x, yw - into * 1.6)
        else:
            xw = 0.0 if d.side == "W" else p.width
            into = 1 if d.side == "W" else -1
            for sgn in (-1, 1):
                hinge = P(xw, cy + sgn * 3.0)
                leaf_end = P(xw + into * 3.0, cy + sgn * 3.0)
                cv.line(hinge, leaf_end, width=0.8, eid=g)
                cv.arc(leaf_end, P(xw + into * 3.0, cy), P(xw, cy), eid=g)
            tx, ty = P(xw + into * 5.5, cy)
        if not d.tagged:
            continue                                   # drafting omission: swing drawn, tag missing
        box = pymupdf.Rect(tx - 13, ty - 6, tx + 13, ty + 6)
        cv.rect(box, width=0.6, eid=("door_tag", d.mark))
        cv.rec(g, box)
        cv.text(tx, ty + 2.5, label, size=7, align="center", eid=("door_tag", d.mark))


def draw_room_labels(cv: Canvas, p: Project, light=False):
    col = GREY if light else BLACK
    for r in p.rooms:
        if r.side == "C":
            cx, cy = P(p.width * 0.5, p.corridor_y)
            cx += 0
        else:
            cx, cy = P((r.x0 + r.x1) / 2, (r.y0 + r.y1) / 2)
        name = p.room_name_override.get(r.number, r.name) if not light else r.name
        if r.side == "C":
            cv.text(cx, cy - 1, name, size=6.5 if not light else 5.5, color=col, align="center", eid=("room_name", r.number))
            cv.text(cx, cy + 7, r.number, size=6 if not light else 5, color=col, align="center")
        else:
            cv.text(cx, cy - 2, name, size=7 if not light else 5.5, color=col, align="center", eid=("room_name", r.number))
            cv.text(cx, cy + 7, r.number, size=6.5 if not light else 5, color=col, align="center")


def callout(cv: Canvas, x, y, num, sheet, target):
    cv.circle((x, y), 14, eid=("callout", sheet))
    cv.line((x - 14, y), (x + 14, y), width=0.5)
    cv.text(x, y - 3, num, size=7, bold=True, align="center", eid=("callout", sheet))
    cv.text(x, y + 9, sheet, size=6, align="center", eid=("callout", sheet))
    cv.line((x + 10, y - 10), target, width=0.35, dashes="[2 2] 0")


def draw_table(cv: Canvas, x0, y0, headers, widths, rows, row_ids, title, row_h=14.0, size=6.5, cell_eids=None):
    cv.text(x0, y0 - 8, title, size=10, bold=True)
    xs = [x0]
    for w in widths:
        xs.append(xs[-1] + w)
    all_rows = [headers] + rows
    for ri, row in enumerate(all_rows):
        y = y0 + ri * row_h
        rid = row_ids[ri - 1] if ri else None
        for ci, val in enumerate(row):
            cell = pymupdf.Rect(xs[ci], y, xs[ci + 1], y + row_h)
            cv.rect(cell, width=0.8 if ri == 0 else 0.4)
            cv.rec(rid, cell)
            if cell_eids and ri:
                cv.rec((cell_eids[0], rid[1], headers[ci]), cell)
            if val:
                cv.text(xs[ci] + 3, y + row_h - 4, str(val), size=size, bold=(ri == 0))
    return xs, y0, row_h


# ----------------------------------------------------------------------------------------------
def sheet_G001(cv: Canvas, p: Project):
    cv.text(60, 110, PROJECT_NAME, size=26, bold=True)
    cv.text(60, 134, "CONSTRUCTION DOCUMENTS - SYNTHETIC SAMPLE SET (FICTIONAL PROJECT, NOT FOR CONSTRUCTION)", size=9)
    rows = [[s, t] for s, t in p.sheet_index]
    draw_table(cv, 60, 190, ["SHEET", "SHEET TITLE"], [70, 260], rows, [("index_row", s) for s, _ in p.sheet_index],
               "SHEET INDEX", row_h=16, size=7.5)
    cv.rec(("index_table", ""), pymupdf.Rect(60, 190, 390, 190 + 16 * (len(rows) + 1)))
    ab = ["AFF  ABOVE FINISHED FLOOR", "EQ  EQUAL", "HM  HOLLOW METAL", "NIC  NOT IN CONTRACT",
          "SIM  SIMILAR", "TYP  TYPICAL", "WD  WOOD", "MCA  MINIMUM CIRCUIT AMPACITY", "V/PH  VOLTS / PHASE"]
    cv.text(470, 182, "ABBREVIATIONS", size=10, bold=True)
    for i, a in enumerate(ab):
        cv.text(470, 200 + 12 * i, a, size=7)
    notes = ["1. ALL DOCUMENTS IN THIS SET ARE SYNTHETIC AND FOR SOFTWARE DEMONSTRATION ONLY.",
             "2. DIMENSIONS GOVERN. DO NOT SCALE DRAWINGS.",
             "3. REFER TO THE PROJECT MANUAL FOR SPECIFICATIONS."]
    cv.text(470, 340, "GENERAL NOTES", size=10, bold=True)
    for i, n in enumerate(notes):
        cv.text(470, 358 + 12 * i, n, size=7)


def sheet_A101(cv: Canvas, p: Project):
    cv.text(40, 34, "GENERAL NOTES:", size=7, bold=True)
    cv.text(40, 45, "1. DIMENSIONS ARE TO GRID LINES UNLESS NOTED OTHERWISE. DO NOT SCALE DRAWINGS.", size=6.5)
    cv.text(40, 55, "2. SEE A-601 FOR DOOR SCHEDULE AND DOOR TYPES. DOOR TAGS SHOW DOOR NUMBER.", size=6.5)
    draw_grids(cv, p)
    draw_dims(cv, p, "A-101")
    draw_building(cv, p)
    draw_doors(cv, p)
    draw_room_labels(cv, p)
    xs, ys = grid_positions(p)
    entry = P(0.5, p.corridor_y + 3.0)
    for i, (num, sheet) in enumerate(p.callouts):
        callout(cv, OX + 40 + 70 * i, ys[-1] + 42, num, sheet, entry)
    cv.text(OX, ys[-1] + 82, "1   FIRST FLOOR PLAN", size=10, bold=True)
    cv.text(OX + 140, ys[-1] + 82, f"SCALE: {scale_label(p)}", size=7)


def sheet_A601(cv: Canvas, p: Project):
    headers = ["MARK", "TYPE", "WIDTH", "HEIGHT", "MATERIAL", "FRAME", "FIRE RATING", "HDW SET", "SPEC SECTION", "REMARKS"]
    widths = [44, 38, 48, 48, 58, 46, 64, 54, 70, 150]
    doors = sorted((d for d in p.doors if d.in_schedule), key=lambda d: d.mark)
    rows = [[d.mark, d.type, d.width, d.height, d.material, d.frame, d.fire, d.hw, d.spec, d.remarks] for d in doors]
    draw_table(cv, 40, 80, headers, widths, rows, [("sched_row", d.mark) for d in doors], "DOOR SCHEDULE",
               cell_eids=("sched_cell",))
    ytypes = 80 + 14 * (len(rows) + 1) + 60
    cv.circle((54, ytypes), 12)
    cv.text(54, ytypes + 3, "4", size=8, bold=True, align="center")
    cv.text(72, ytypes + 3, "DOOR TYPES", size=10, bold=True)
    for i, (t, info) in enumerate(DOOR_TYPES.items()):
        x = 80 + 150 * i
        w = 36 if t != "D3" else 72
        cv.rect(pymupdf.Rect(x, ytypes + 24, x + w, ytypes + 108), width=0.8)
        cv.text(x, ytypes + 122, f"{t} - {info['desc']}", size=6.5, bold=True)
        cv.text(x, ytypes + 132, f"SEE SECTION {info['spec']}", size=6)
    cv.text(560, ytypes + 30, "NOTES:", size=7, bold=True)
    cv.text(560, ytypes + 41, "1. HARDWARE SETS PER SECTION 08 71 00.", size=6.5)
    cv.text(560, ytypes + 51, "2. - IN FIRE RATING COLUMN = NON-RATED OPENING.", size=6.5)


def sheet_S101(cv: Canvas, p: Project):
    cv.text(40, 34, "GENERAL NOTES:", size=7, bold=True)
    cv.text(40, 45, "1. RTU SUPPORT FRAMES ARE DESIGNED FOR THE MAX OPERATING WEIGHTS NOTED. COORDINATE WITH M-601.", size=6.5)
    cv.text(40, 55, "2. DO NOT SCALE DRAWINGS.", size=6.5)
    draw_grids(cv, p)
    draw_dims(cv, p, "S-101")
    xs, ys = grid_positions(p)
    cv.rect(pymupdf.Rect(xs[0], ys[0], xs[-1], ys[-1]), width=0.5, color=GREY)
    for gx in xs:
        for gy in ys:
            cv.rect(pymupdf.Rect(gx - 3, gy - 3, gx + 3, gy + 3), width=0.5, fill=BLACK)
    for gy in ys:
        cv.line((xs[0], gy), (xs[-1], gy), width=1.0)
    for e in p.equips:
        if not e.tag.startswith("RTU") or not e.plan:
            continue
        x, y, _ = e.plan[0]
        fx, fy = P(x, y)
        frame = pymupdf.Rect(fx - 30, fy - 18, fx + 30, fy + 18)
        cv.rect(frame, width=0.8, dashes="[4 2] 0")
        cv.text(fx, fy - 2, f"{e.tag} SUPPORT FRAME", size=6, align="center")
        cv.text(fx, fy + 30, f"{e.tag}: DESIGN MAX OPERATING WEIGHT {p.rtu_design_max[e.tag]:,} LB", size=6,
                align="center", eid=("rtu_note", e.tag))
    cv.text(OX, ys[-1] + 82, "1   ROOF FRAMING PLAN", size=10, bold=True)
    cv.text(OX + 150, ys[-1] + 82, f"SCALE: {scale_label(p)}", size=7)


def sheet_M101(cv: Canvas, p: Project):
    cv.text(40, 34, "GENERAL NOTES:", size=7, bold=True)
    cv.text(40, 45, "1. RTUs ARE ROOF-MOUNTED ABOVE; SEE S-101 FOR SUPPORT FRAMING AND M-601 FOR SCHEDULE.", size=6.5)
    cv.text(40, 55, "2. EQUIPMENT MARKED (NIC) IS FURNISHED AND INSTALLED BY OWNER.", size=6.5)
    draw_grids(cv, p)
    draw_building(cv, p, light=True, doors=False)
    draw_room_labels(cv, p, light=True)
    for e in p.equips:
        for (x, y, qual) in e.plan:
            ex, ey = P(x, y)
            if e.tag.startswith("RTU"):
                cv.rect(pymupdf.Rect(ex - 27, ey - 16, ex + 27, ey + 16), width=0.8, dashes="[4 2] 0")
            else:
                cv.rect(pymupdf.Rect(ex - 5, ey - 5, ex + 5, ey + 5), width=0.6)
            box = pymupdf.Rect(ex - 18, ey - 30, ex + 18, ey - 19)
            cv.rect(box, width=0.6, eid=("equip_tag", e.tag))
            cv.text(ex, ey - 21.5, e.tag, size=6.5, align="center", eid=("equip_tag", e.tag))
            if qual:
                cv.text(ex + 21, ey - 21.5, qual, size=6, eid=("equip_tag", e.tag))
    xs, ys = grid_positions(p)
    cv.text(OX, ys[-1] + 82, "1   MECHANICAL FLOOR PLAN", size=10, bold=True)
    cv.text(OX + 180, ys[-1] + 82, f"SCALE: {scale_label(p)}", size=7)


def sheet_M601(cv: Canvas, p: Project):
    headers = ["TAG", "DESCRIPTION", "QTY", "CAPACITY", "V/PH", "MCA (A)", "WEIGHT (LB)", "SPEC", "REMARKS"]
    widths = [50, 130, 32, 62, 44, 48, 62, 56, 150]
    eq = [e for e in p.equips if e.in_schedule]
    rows = [[e.tag, e.desc, str(e.qty), e.capacity, e.volts, str(e.mca) if e.mca else "-", f"{e.weight:,}",
             e.spec, e.remarks] for e in eq]
    draw_table(cv, 40, 80, headers, widths, rows, [("sched_row", e.tag) for e in eq], "MECHANICAL EQUIPMENT SCHEDULE",
               cell_eids=("sched_cell",))
    y = 80 + 14 * (len(rows) + 1) + 30
    cv.text(40, y, "NOTES:", size=7, bold=True)
    cv.text(40, y + 11, "1. ELECTRICAL CHARACTERISTICS SHALL BE COORDINATED WITH E-601 PRIOR TO ORDERING.", size=6.5)
    cv.text(40, y + 21, "2. RTU OPERATING WEIGHTS SHALL NOT EXCEED STRUCTURAL DESIGN WEIGHTS ON S-101.", size=6.5)


def sheet_E601(cv: Canvas, p: Project):
    headers = ["PANEL", "CKT", "LOAD", "V/PH", "POLES", "BKR (A)"]
    widths = [44, 60, 170, 50, 44, 50]
    rows, ids = [], []
    for e in p.equips:
        if e.circuit:
            c = e.circuit
            rows.append([c.panel, c.ckt, e.tag, c.volts, str(c.poles), str(c.bkr)])
            ids.append(("panel_row", e.tag))
    rows += [["LA", "13", "LIGHTING - CORRIDOR", "120/1", "1", "20"], ["LA", "15", "RECEPTACLES - OFFICES", "120/1", "1", "20"]]
    ids += [("panel_row", "LTG"), ("panel_row", "REC")]
    order = sorted(range(len(rows)), key=lambda i: (rows[i][0], int(rows[i][1].split(",")[0])))
    rows = [rows[i] for i in order]
    ids = [ids[i] for i in order]
    cv.text(40, 50, "PANEL HA: 480Y/277V, 3PH, 4W    PANEL LA: 208Y/120V, 3PH, 4W", size=7)
    draw_table(cv, 40, 80, headers, widths, rows, ids, "PANEL SCHEDULES (BRANCH CIRCUITS)")


SHEET_FUNCS = {"G-001": sheet_G001, "A-101": sheet_A101, "A-601": sheet_A601, "S-101": sheet_S101,
               "M-101": sheet_M101, "M-601": sheet_M601, "E-601": sheet_E601}
SHEET_SCALE = {"A-101": "1/8\" = 1'-0\"", "S-101": "1/8\" = 1'-0\"", "M-101": "1/8\" = 1'-0\""}


def render_set(p: Project, path: Path) -> dict:
    """Render one revision of the drawing set. Returns {sheet: {"page": i, "elements": {...}}}."""
    global SCALE
    SCALE = SCALE_METRIC if p.units == "metric" else SCALE_IMPERIAL
    doc = pymupdf.open()
    index = {}
    for i, (sheet, title) in enumerate(SHEETS):
        page = doc.new_page(width=W, height=H)
        cv = Canvas(page)
        page.draw_rect(pymupdf.Rect(10, 10, W - 10, H - 10), color=BLACK, width=1.5)
        SHEET_FUNCS[sheet](cv, p)
        title_block(cv, p, sheet, title, scale_label(p) if sheet in SHEET_SCALE else "AS NOTED")
        for target in p.clouds.get(sheet, []):
            if target in cv.elements:
                tight = target[0] in ("sched_row", "sched_cell")     # table rows: hug the row
                cv.cloud(cv.elements[target], p.rev, margin=1.5 if tight else 6.0, scallop=6.0 if tight else 9.0)
        index[sheet] = {"page": i, "elements": cv.elements}
    doc.set_metadata({"title": f"{PROJECT_NAME} - CD Set Rev {p.rev} (synthetic)", "author": FIRM_NAME,
                      "creator": "aecdi synthetic generator"})
    doc.save(path, garbage=3, deflate=True)
    return index


# ----------------------------------------------------------------------------------------------
def spec_text(p: Project) -> dict[str, tuple[str, list[tuple[str, str]]]]:
    hw = [("A.", "HW-1: Office lockset, three hinges, wall stop."),
          ("B.", "HW-2: Storeroom lockset, closer, three hinges, kick plate."),
          ("C.", "HW-3: Privacy set with occupancy indicator, closer, three hinges, coat hook."),
          ("D.", "HW-4: Exit devices (pair), closers, continuous hinges, threshold and weatherstripping.")]
    return {
        "08 11 13": ("HOLLOW METAL DOORS AND FRAMES", [
            ("PART 1", "GENERAL"), ("1.1", "SUMMARY"),
            ("A.", "Section includes hollow metal doors and frames."),
            ("B.", "Related sections: Section 08 71 00 \"Door Hardware\"."),
            ("PART 2", "PRODUCTS"), ("2.1", "HOLLOW METAL DOORS"),
            ("A.", "Interior doors: SDI A250.8 Level 2, 18 gauge, 1-3/4 inches thick, seamless edges."),
            ("B.", "Fire-rated doors: labeled assemblies rated 20, 45, 60 and 90 minutes and 3 hours as indicated."),
            ("2.2", "FRAMES"), ("A.", "Frames: 16 gauge, welded, with three silencers per strike jamb."),
            ("PART 3", "EXECUTION"), ("3.1", "INSTALLATION"),
            ("A.", "Install doors and frames plumb, rigid and in true alignment."),
        ]),
        "08 14 16": ("FLUSH WOOD DOORS", [
            ("PART 1", "GENERAL"), ("1.1", "SUMMARY"),
            ("A.", "Section includes solid-core flush wood doors with wood veneer faces."),
            ("PART 2", "PRODUCTS"), ("2.1", "FLUSH WOOD DOORS"),
            ("A.", "Doors: solid core flush wood doors, 1-3/4 inches thick, WDMA I.S.1A Extra Heavy Duty."),
            ("B.", "Faces: rotary-cut red oak veneer, Grade A, for transparent finish."),
            ("C.", "Fire-rated doors: provide wood doors rated 20, 45 and 60 minutes where indicated. "
                   "Wood doors rated more than 60 minutes are not permitted; use hollow metal doors."),
            ("D.", "Hollow-core doors are not permitted."),
            ("PART 3", "EXECUTION"), ("3.1", "INSTALLATION"),
            ("A.", "Install doors in frames with uniform clearances of 1/8 inch at heads and jambs."),
        ]),
        "08 41 13": ("ALUMINUM-FRAMED ENTRANCES AND STOREFRONTS", [
            ("PART 2", "PRODUCTS"), ("2.1", "ENTRANCE DOORS"),
            ("A.", "Entrance doors: medium stile aluminum, 1-3/4 inches thick, pairs as indicated, insulated glazing."),
        ]),
        "08 71 00": ("DOOR HARDWARE", [
            ("PART 1", "GENERAL"), ("1.1", "SUMMARY"),
            ("A.", "Section includes door hardware for wood, hollow metal and aluminum doors."),
            ("PART 3", "EXECUTION"), ("3.5", "HARDWARE SETS"), *hw,
        ]),
        "23 34 16": ("CENTRIFUGAL HVAC FANS", [
            ("PART 2", "PRODUCTS"), ("2.1", "CEILING-MOUNTED EXHAUST FANS"),
            ("A.", "Fans: direct drive, 120 V, 1 phase, 60 Hz, with backdraft damper and speed controller."),
        ]),
        "23 36 00": ("AIR TERMINAL UNITS", [
            ("PART 2", "PRODUCTS"), ("2.1", "SINGLE-DUCT VAV TERMINALS"),
            ("A.", "Terminals: pressure independent, with DDC controls furnished under Division 23."),
        ]),
        "23 74 13": ("PACKAGED OUTDOOR CENTRAL-STATION AIR-HANDLING UNITS", [
            ("PART 2", "PRODUCTS"), ("2.1", "ROOFTOP UNITS"),
            ("A.", "Units: factory-assembled rooftop units with gas heat and DX cooling, sized as scheduled."),
            ("B.", f"Refrigerant: {SPEC_REFRIGERANT}. R-410A is not permitted."),
            ("C.", "Electrical characteristics: 460 V, 3 phase, 60 Hz, single-point power connection."),
            ("D.", "Economizer: integrated, with barometric relief."),
        ]),
    }


def _wrap(text: str, width_pt: float, size: float, font="helv") -> list[str]:
    words, lines, cur = text.split(), [], ""
    for wd in words:
        trial = f"{cur} {wd}".strip()
        if pymupdf.get_text_length(trial, fontname=font, fontsize=size) > width_pt and cur:
            lines.append(cur)
            cur = wd
        else:
            cur = trial
    if cur:
        lines.append(cur)
    return lines


def render_spec(p: Project, path: Path):
    doc = pymupdf.open()
    content = spec_text(p)
    for code in sorted(p.spec_sections):
        title, paras = content[code]
        page = doc.new_page(width=612, height=792)
        page.insert_text((54, 40), f"{PROJECT_NAME} - PROJECT MANUAL (SYNTHETIC SAMPLE)", fontsize=7, fontname="helv",
                         color=GREY)
        page.insert_text((54, 90), f"SECTION {code}", fontsize=11, fontname="hebo")
        page.insert_text((54, 106), title, fontsize=11, fontname="hebo")
        y = 136
        for label, text in paras:
            if label.startswith("PART"):
                y += 6
                page.insert_text((54, y), f"{label} - {text}", fontsize=9, fontname="hebo")
                y += 16
            elif label[0].isdigit():
                page.insert_text((54, y), f"{label}   {text}", fontsize=9, fontname="hebo")
                y += 14
            else:
                lines = _wrap(text, 440, 9)
                page.insert_text((72, y), label, fontsize=9, fontname="helv")
                for ln in lines:
                    page.insert_text((90, y), ln, fontsize=9, fontname="helv")
                    y += 12
                y += 4
        page.insert_text((54, y + 20), f"END OF SECTION {code}", fontsize=9, fontname="hebo")
        page.insert_text((250, 760), f"{title[:40]}   {code} - 1", fontsize=7, fontname="helv", color=GREY)
    doc.set_metadata({"title": f"{PROJECT_NAME} - Project Manual (synthetic)", "author": FIRM_NAME})
    doc.save(path, garbage=3, deflate=True)


def render_boq(p: Project, path: Path) -> dict:
    wb = Workbook()
    ws = wb.active
    ws.title = "BOQ"
    ws["A1"] = f"{PROJECT_NAME} - BILL OF QUANTITIES (REV {p.rev})"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = "Synthetic sample for software demonstration. Quantities are for pricing only."
    ws.merge_cells("A1:G1")
    headers = ["ITEM", "DESCRIPTION", "SPEC SECTION", "UNIT", "QTY", "RATE (USD)", "AMOUNT (USD)"]
    ws.append([])
    ws.append(headers)
    for c in ws[4]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDDDDD")
    rows: dict[str, int] = {}
    div = None
    for line in p.boq:
        d = line.item[:2]
        if d != div:
            ws.append([f"DIVISION {d} - {'OPENINGS' if d == '08' else 'HVAC'}"])
            ws.cell(ws.max_row, 1).font = Font(bold=True, italic=True)
            div = d
        ws.append([line.item, line.desc, line.spec, line.unit, line.qty, line.rate, None])
        r = ws.max_row
        ws.cell(r, 7, f"=E{r}*F{r}")
        rows[line.key] = r
    ws.append([])
    ws.append(["", "TOTAL", "", "", "", "", f"=SUM(G5:G{ws.max_row - 1})"])
    for col, w in zip("ABCDEFG", [10, 58, 14, 7, 7, 12, 14]):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=5):
        for c in row:
            c.alignment = Alignment(vertical="top")
    wb.save(path)
    return rows


NATIVE_COLUMNS = [("Mark", "mark"), ("Type Mark", "type"), ("Width", "width"), ("Height", "height"),
                  ("Door Material", "material"), ("Frame Material", "frame"), ("Fire Rating", "fire"),
                  ("Hardware Set", "hw"), ("Spec Section", "spec")]


def _revit_len(v: str) -> str:
    """Revit's export format for lengths: 3' - 0"."""
    return v.replace("'-", "' - ") if "'" in v else v


def render_native_schedule(p: Project, path: Path) -> dict:
    """Door schedule exported from the Revit model (Report > Schedule, tab-delimited .txt).
    Returns {mark: line number}. model_overrides holds values where the model differs from the issued PDF."""
    lines = ["\"Door Schedule\"", "\t".join(f'"{h}"' for h, _ in NATIVE_COLUMNS), ""]
    rows = {}
    for d in sorted((d for d in p.doors if d.in_schedule), key=lambda d: d.mark):
        vals = []
        for h, attr in NATIVE_COLUMNS:
            v = p.model_overrides.get(d.mark, {}).get(h, getattr(d, attr))
            vals.append(_revit_len(v) if attr in ("width", "height") else v)
        lines.append("\t".join('"' + v.replace('"', '""') + '"' for v in vals))   # CSV-style quoting
        rows[d.mark] = len(lines)
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return rows


def render_redlines(rev_a_pdf: Path, p: Project, idx_a: dict, idx_b: dict, path: Path) -> list[dict]:
    """Client review comments marked up on the Rev A set (standard PDF annotations, as from Bluebeam)."""
    doc = pymupdf.open(rev_a_pdf)
    placed = []
    for m in p.markups:
        idx = idx_b if m["from"] == "b" else idx_a
        el = idx[m["sheet"]]["elements"].get(tuple(m["element"]))
        page_no = idx_a[m["sheet"]]["page"]
        r = pymupdf.Rect(el) + (-8, -8, 8, 8)
        a = doc[page_no].add_rect_annot(r)
        a.set_colors(stroke=(0.85, 0.05, 0.05))
        a.set_border(width=1.5)
        a.set_info(title="Client reviewer", content=m["text"], subject="Redline")
        a.update()
        placed.append({"sheet": m["sheet"], "page": page_no, "bbox": [round(v, 1) for v in r], "text": m["text"]})
    doc.save(path, garbage=3, deflate=True)
    return placed


def generate(seed: int, out_dir: Path, clean: bool = False, units: str = "imperial") -> Path:
    a, b = make_project(seed, clean, units)
    prefix = "clean" if clean else ("metric" if units == "metric" else "seed")
    out = out_dir / f"{prefix}_{seed:02d}"
    out.mkdir(parents=True, exist_ok=True)
    idx_a = render_set(a, out / "CD_Set_RevA.pdf")
    idx_b = render_set(b, out / "CD_Set_RevB.pdf")
    render_spec(b, out / "Project_Manual.pdf")
    boq_rows = render_boq(b, out / "BOQ_RevB.xlsx")
    render_native_schedule(b, out / "Revit_Door_Schedule_Export.txt")
    redlines = render_redlines(out / "CD_Set_RevA.pdf", b, idx_a, idx_b, out / "CD_Set_RevA_REDLINES.pdf")

    def resolve(g):
        if "redline" in g:
            r = redlines[g["redline"]]
            return {"doc": "CD_Set_RevA_REDLINES.pdf", "sheet": r["sheet"], "page": r["page"], "bbox": r["bbox"]}
        if g["sheet"] == "BOQ":
            key = g["element"][1]
            return {"doc": "BOQ_RevB.xlsx", "xlsx_sheet": "BOQ", "row": boq_rows.get(key)}
        sheet = g["sheet"]
        el = idx_b[sheet]["elements"].get(tuple(g["element"])) or idx_a[sheet]["elements"].get(tuple(g["element"]))
        return {"doc": "CD_Set_RevB.pdf", "sheet": sheet, "page": idx_b[sheet]["page"],
                "bbox": [round(v, 1) for v in el] if el is not None else None}

    key = {
        "project": PROJECT_NAME, "seed": seed, "clean": clean, "units": units,
        "files": {"rev_a": "CD_Set_RevA.pdf", "rev_b": "CD_Set_RevB.pdf", "spec": "Project_Manual.pdf",
                  "boq": "BOQ_RevB.xlsx", "redlines": "CD_Set_RevA_REDLINES.pdf",
                  "native_schedule": "Revit_Door_Schedule_Export.txt"},
        "defects": [{"id": d.id, "kind": d.kind, "description": d.description, "severity": d.severity,
                     "accept": d.accept, "gt": [resolve(g) for g in d.gt]} for d in b.defects],
        "traps": b.traps,
        "legit_revision_changes": ["A-101: new door (clouded)", "A-601: new schedule row (clouded)",
                                   "M-601: RTU-1 capacity revised (clouded)"],
        "redlines": redlines,
    }
    (out / "answer_key.json").write_text(json.dumps(key, indent=2), encoding="utf-8")
    return out
