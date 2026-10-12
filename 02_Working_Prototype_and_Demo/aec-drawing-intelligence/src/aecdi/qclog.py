"""Score the tool against a real QC log (pilot evaluation on Uppteam's past projects).

    aecdi score-qclog <out_dir>/findings.json <qc_log.csv> --out qc_match.xlsx

QC log CSV columns (template: eval_templates/qc_log_template.csv):
    issue_id, sheets, category, severity, description, covered
`sheets` is a list like "A-101; A-601". `covered` = yes/no: whether the issue type is one the tool
attempts (see the coverage matrix in Deliverable 03); recall is reported for covered issues only.

Matching is a *suggestion* for a human: an issue and a finding are candidates when they share a sheet
and at least one identifier (door mark, equipment tag, grid span, sheet number, spec section, BOQ
item). The reviewer confirms each match in the workbook; the summary sheet recomputes from those
confirmations with Excel formulas.
"""
from __future__ import annotations

import csv
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

TOKEN_RES = [
    re.compile(r"\b(?:RTU|AHU|EF|SF|VAV|FCU|CUH|EUH|UH|CU|HP|ERV|DOAS)-\d+[A-Z]?\b"),
    re.compile(r"\b[A-Z]{1,2}-\d{3}\b"),                 # sheet numbers
    re.compile(r"\b\d{2} \d{2} \d{2}\b"),                # spec sections
    re.compile(r"\b\d{2}-\d{3}\b"),                      # BOQ items
    re.compile(r"\bHW-\d+\b"),
    re.compile(r"\b[DW]\d{1,2}\b"),                      # door / window types
    re.compile(r"\b(?:grid\s*)?(\d{1,2}|[A-H])-(\d{1,2}|[A-H])\b", re.I),
    re.compile(r"\b\d{3}[A-Z]?\b"),                      # door / room numbers
]
CATEGORY_WORDS = {
    "missing_info": ("missing", "not defined", "blank", "no schedule", "not in", "undefined", "untagged",
                     "does not appear", "not shown", "not tagged", "no row", "no tag"),
    "quantity_mismatch": ("quantity", "count", "boq", "qty", "priced", "nic"),
    "dimension_conflict": ("dimension", "grid", "chain", "overall"),
    "cross_discipline_conflict": ("voltage", "breaker", "circuit", "mca", "weight", "structural", "electrical"),
    "spec_conflict": ("spec", "section", "rating", "material", "refrigerant"),
    "revision_issue": ("cloud", "revision", "addendum", "markup", "redline", "not incorporated"),
    "tag_conflict": ("duplicate", "twice"),
    "model_mismatch": ("model", "revit", "out of date"),
}


def tokens(text: str) -> set[str]:
    out = set()
    for rx in TOKEN_RES:
        for m in rx.finditer(text or ""):
            t = m.group(0).upper().replace("GRID", "").strip()
            out.add(t)
    return out


def finding_tokens(f: dict) -> set[str]:
    t = tokens(f["statement"])
    for e in f["entities"]:
        t |= tokens(e.split(":", 1)[-1])
    return t


def guess_category(text: str) -> str | None:
    low = (text or "").lower()
    scores = {k: sum(w in low for w in ws) for k, ws in CATEGORY_WORDS.items()}
    k = max(scores, key=scores.get)
    return k if scores[k] else None


def load_qc_log(path: Path) -> list[dict]:
    with open(path, newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    for r in rows:
        r["sheet_list"] = {s.strip().upper() for s in re.split(r"[;,]", r.get("sheets", "")) if s.strip()}
        r["covered"] = str(r.get("covered", "yes")).strip().lower() not in ("no", "n", "false", "0")
    return rows


def match(findings: list[dict], issues: list[dict]) -> list[dict]:
    shipped = [f for f in findings if f["tier"] != "needs_review"]
    out = []
    for iss in issues:
        it = tokens(iss.get("description", ""))
        cat = iss.get("category") or guess_category(iss.get("description", ""))
        best, best_score = None, 0.0
        for f in shipped:
            same_sheet = not iss["sheet_list"] or bool(iss["sheet_list"] & {s.upper() for s in f["sheets"]})
            shared = it & finding_tokens(f)
            if not same_sheet or not shared:
                continue
            score = len(shared) + (1.0 if cat and cat == f["type"] else 0.0)
            if score > best_score:
                best, best_score = f, score
        out.append({"issue": iss, "finding": best, "score": best_score})
    return out


def score_qc_log(findings_path: Path, qc_path: Path, out_xlsx: Path) -> dict:
    result = json.loads(Path(findings_path).read_text(encoding="utf-8"))
    findings = result["findings"]
    issues = load_qc_log(qc_path)
    matches = match(findings, issues)
    wb = Workbook()
    ws = wb.active
    ws.title = "QC issues vs tool"
    hdr = ["QC issue", "Sheets", "Category", "Severity", "Covered", "QC description", "Suggested finding", "Finding",
           "Match score", "Reviewer confirms match (Yes/No)"]
    ws.append(hdr)
    for m in matches:
        i, f = m["issue"], m["finding"]
        ws.append([i.get("issue_id"), i.get("sheets"), i.get("category"), i.get("severity"), "yes" if i["covered"] else "no",
                   i.get("description"), f["id"] if f else "", f["statement"] if f else "", round(m["score"], 1),
                   "Yes" if f and m["score"] >= 2 else ""])
    n = ws.max_row
    dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"J2:J{max(2, n)}")
    matched_ids = {m["finding"]["id"] for m in matches if m["finding"]}
    ws2 = wb.create_sheet("Tool findings not in QC log")
    ws2.append(["Finding", "Tier", "Severity", "Type", "Statement", "Reviewer verdict (new real issue / false positive)"])
    for f in findings:
        if f["tier"] != "needs_review" and f["id"] not in matched_ids:
            ws2.append([f["id"], f["tier"], f["severity"], f["type"], f["statement"], ""])
    ws3 = wb.create_sheet("Summary")
    ws3.append(["Metric", "Value", "How it is computed"])
    ws3.append(["Covered QC issues", f'=COUNTIF(\'QC issues vs tool\'!E2:E{n},"yes")', "Issue types the tool attempts"])
    ws3.append(["Covered issues found (confirmed)",
                f'=COUNTIFS(\'QC issues vs tool\'!E2:E{n},"yes",\'QC issues vs tool\'!J2:J{n},"Yes")', "Reviewer-confirmed matches"])
    ws3.append(["Recall on covered issues", "=IFERROR(B3/B2,0)", "Pilot gate: >= 80%"])
    ws3.append(["Tool findings (shipped)", sum(1 for f in findings if f["tier"] != "needs_review"), ""])
    ws3.append(["Findings matched to QC log", f'=COUNTIF(\'QC issues vs tool\'!J2:J{n},"Yes")', ""])
    ws3.append(["Unmatched findings judged real", "", "Fill in from the second sheet: real issues the QC missed"])
    ws3.append(["Precision", "=IFERROR((B6+B7)/B5,0)", "Pilot gate: >= 85%"])
    for w_, c in zip((8, 12, 18, 9, 9, 60, 10, 60, 8, 16), "ABCDEFGHIJ"):
        ws.column_dimensions[c].width = w_
    for sh in (ws, ws2, ws3):
        for c in sh[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3A5F")
    wb.save(out_xlsx)
    covered = [m for m in matches if m["issue"]["covered"]]
    suggested = sum(1 for m in covered if m["finding"] and m["score"] >= 2)
    return {"qc_issues": len(issues), "covered": len(covered), "suggested_matches": suggested,
            "provisional_recall": round(suggested / len(covered), 3) if covered else None,
            "unmatched_findings": ws2.max_row - 1, "workbook": str(out_xlsx)}
