"""Outputs that drop into an existing AEC QC workflow.

* findings.json            machine-readable result (review UI, integrations)
* QC_Issue_Log.xlsx        one row per finding, Excel QC-log style, with evidence and review columns
* *_REVIEW.pdf             copies of the sources with a rectangle markup + comment per evidence anchor
                           (standard PDF annotations: open in Bluebeam Revu / Acrobat)
* acc_issues.json          Autodesk Construction Cloud Issues payloads (dry run; posting needs a project
                           id + OAuth token and is done by the n8n workflow after human acceptance)
* project_graph.json       nodes + typed edges + provenance
* revision_delta.json      every Rev A -> Rev B change with its clouded flag
* overlays/*.png           red = only in previous rev, blue = only in current rev
"""
from __future__ import annotations

import json
import shutil
from collections import defaultdict
from pathlib import Path

import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from .graph import to_json
from .models import Finding

SEV_COLOR = {"high": (0.86, 0.1, 0.1), "medium": (0.95, 0.55, 0.0), "low": (0.1, 0.45, 0.85)}
TIER_FILL = {"confirmed": "F8D7DA", "probable": "FFF3CD", "needs_review": "E2E3E5"}


def write_all(o, out: Path):
    out.mkdir(parents=True, exist_ok=True)
    r = o.result
    (out / "findings.json").write_text(r.model_dump_json(indent=2), encoding="utf-8")
    (out / "project_graph.json").write_text(json.dumps(to_json(o.graph), indent=2), encoding="utf-8")
    (out / "revision_delta.json").write_text(json.dumps([c.model_dump() for c in r.revision_changes], indent=2),
                                             encoding="utf-8")
    if o.llm_log:
        (out / "llm_log.json").write_text(json.dumps(o.llm_log, indent=2), encoding="utf-8")
    issue_log(r.findings, out / "QC_Issue_Log.xlsx", r.stats)
    files = {p.name: p for p in (o.inputs.rev_b, o.inputs.rev_a, o.inputs.spec) if p}
    annotate(r.findings, files, out)
    (out / "acc_issues.json").write_text(json.dumps(acc_payloads(r.findings), indent=2), encoding="utf-8")
    # copy inputs so the review UI can render pages from the out folder alone
    src = out / "sources"
    src.mkdir(exist_ok=True)
    for p in (o.inputs.rev_b, o.inputs.rev_a, o.inputs.spec, o.inputs.boq):
        if p:
            shutil.copy2(p, src / p.name)


def _ev_text(f: Finding) -> str:
    parts = []
    for e in f.evidence:
        loc = f"{e.xlsx_sheet}!{e.cell}" if e.xlsx_sheet else f"{e.sheet} p{(e.page or 0) + 1}"
        parts.append(f"[{e.doc} {loc}] {e.role}: \"{e.snippet}\"")
    return "\n".join(parts)


def issue_log(findings: list[Finding], path: Path, stats: dict):
    wb = Workbook()
    ws = wb.active
    ws.title = "QC Issues"
    headers = ["ID", "Tier", "Severity", "Type", "Disciplines", "Sheets", "Potential discrepancy", "Evidence",
               "Method", "Confidence", "Why this tier", "Status", "Reviewer", "Reviewer comment"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for f in findings:
        ws.append([f.id, f.tier, f.severity, f.type.replace("_", " "), ", ".join(f.disciplines), ", ".join(f.sheets),
                   f.statement, _ev_text(f), f.method, round(f.confidence, 2), f.rationale, "Open", "", ""])
        r = ws.max_row
        ws.cell(r, 2).fill = PatternFill("solid", fgColor=TIER_FILL[f.tier])
        for c in ws[r]:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    dv = DataValidation(type="list", formula1='"Open,Accepted,Rejected,Deferred"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"L2:L{max(2, ws.max_row)}")
    for col, w in zip("ABCDEFGHIJKLMN", [7, 12, 9, 20, 10, 14, 60, 70, 16, 10, 40, 11, 14, 30]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:N{ws.max_row}"
    s = wb.create_sheet("Run summary")
    s.append(["Metric", "Value"])
    for k, v in stats.items():
        s.append([k, v])
    s.append([])
    s.append(["Note", "All findings are potential discrepancies for engineer review. The tool does not approve, seal or "
                      "modify drawings; the AOR/EOR decides."])
    s.column_dimensions["A"].width = 24
    s.column_dimensions["B"].width = 90
    wb.save(path)


def annotate(findings: list[Finding], files: dict[str, Path], out: Path):
    by_doc = defaultdict(list)
    for f in findings:
        for e in f.evidence:
            if e.page is not None and e.bbox and e.doc in files:
                by_doc[e.doc].append((f, e))
    for doc_name, items in by_doc.items():
        doc = pymupdf.open(files[doc_name])
        for f, e in items:
            page = doc[e.page]
            rect = pymupdf.Rect(e.bbox) + (-3, -3, 3, 3)
            a = page.add_rect_annot(rect)
            a.set_colors(stroke=SEV_COLOR[f.severity])
            a.set_border(width=1.5, dashes=[3, 2] if f.tier != "confirmed" else None)
            a.set_info(title=f"{f.id} [{f.tier}/{f.severity}]",
                       content=f"{f.statement}\n\nEvidence role: {e.role}\nMethod: {f.method}  Confidence: {f.confidence:.2f}",
                       subject=f.type)
            a.update()
            page.insert_text((rect.x0, rect.y0 - 2), f.id, fontsize=6, color=SEV_COLOR[f.severity])
        doc.save(out / f"{Path(doc_name).stem}_REVIEW.pdf", garbage=3, deflate=True)


def acc_payloads(findings: list[Finding]) -> list[dict]:
    """Shape follows the ACC Issues API (POST /construction/issues/v1/projects/{projectId}/issues).
    issueSubtypeId / rootCauseId / projectId are placeholders resolved by the integration at post time."""
    out = []
    for f in findings:
        if f.tier == "needs_review":
            continue
        pin = next((e for e in f.evidence if e.page is not None and e.bbox and e.doc.lower().endswith(".pdf")), None)
        payload = {
            "title": f"[{f.id}] {f.type.replace('_', ' ')}: {', '.join(f.entities)}"[:200],
            "description": f"{f.statement}\n\nEvidence:\n{_ev_text(f)}\n\nGenerated by AEC Drawing Intelligence "
                           f"({f.method}, confidence {f.confidence:.2f}, tier {f.tier}). Requires engineer review.",
            "status": "open",
            "issueSubtypeId": "<resolve: Design / Coordination>",
            "rootCauseId": None,
            "published": False,
            "customAttributes": [],
            "_aecdi": {"finding_id": f.id, "severity": f.severity, "tier": f.tier, "sheets": f.sheets},
        }
        if pin:
            payload["linkedDocuments"] = [{"type": "TwoDVectorPushpin", "urn": f"<resolve: {pin.doc}>",
                                           "details": {"viewable": {"name": pin.sheet},
                                                       "position": {"x": round((pin.bbox[0] + pin.bbox[2]) / 2, 1),
                                                                    "y": round((pin.bbox[1] + pin.bbox[3]) / 2, 1)},
                                                       "_page_index": pin.page}}]
        out.append(payload)
    return out
