"""Bill of quantities ingestion (Excel). Finds the header row instead of assuming row 1,
skips division headings and totals, and keeps the row/cell address of every value."""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ITEM_RE = re.compile(r"^\d{2}-\d{3}[A-Z]?$")
SPEC_RE = re.compile(r"\b(\d{2}\s\d{2}\s\d{2})\b")
DOOR_TYPE_RE = re.compile(r"\btype\s+(D\d+)\b", re.I)
EQUIP_RE = re.compile(r"\b((?:RTU|AHU|EF|SF|VAV|FCU|CUH|EUH|UH|CU|HP|ERV|DOAS)-\d+[A-Z]?)\b")


@dataclass
class BoqItem:
    item: str
    desc: str
    spec: str | None
    unit: str
    qty: float | None
    sheet: str
    row: int
    cols: dict[str, str]        # logical column -> cell address

    @property
    def door_type(self) -> str | None:
        m = DOOR_TYPE_RE.search(self.desc)
        return m.group(1).upper() if m else None

    @property
    def equipment_tag(self) -> str | None:
        m = EQUIP_RE.search(self.desc)
        return m.group(1) if m else None


@dataclass
class Boq:
    path: Path
    items: list[BoqItem]

    @property
    def doc(self) -> str:
        return self.path.name


HEADERS = {"item": ("ITEM",), "desc": ("DESCRIPTION", "DESC"), "spec": ("SPEC",), "unit": ("UNIT",),
           "qty": ("QTY", "QUANTITY")}


def load_boq(path: str | Path) -> Boq:
    path = Path(path)
    wb = load_workbook(path, data_only=False)
    items: list[BoqItem] = []
    for ws in wb.worksheets:
        header_row, cols = None, {}
        for r in range(1, min(ws.max_row, 30) + 1):
            vals = {c: str(ws.cell(r, c).value or "").strip().upper() for c in range(1, ws.max_column + 1)}
            found = {k: next((c for c, v in vals.items() if any(v.startswith(h) for h in hs)), None)
                     for k, hs in HEADERS.items()}
            if found["item"] and found["qty"]:
                header_row, cols = r, found
                break
        if header_row is None:
            continue
        for r in range(header_row + 1, ws.max_row + 1):
            item = str(ws.cell(r, cols["item"]).value or "").strip()
            if not ITEM_RE.match(item):
                continue

            def val(k):
                return ws.cell(r, cols[k]).value if cols.get(k) else None

            qty = val("qty")
            try:
                qty = float(qty) if qty is not None and str(qty).strip() != "" else None
            except ValueError:
                qty = None
            spec_raw = str(val("spec") or "")
            m = SPEC_RE.search(spec_raw)
            items.append(BoqItem(item=item, desc=str(val("desc") or ""), spec=m.group(1) if m else None,
                                 unit=str(val("unit") or ""), qty=qty, sheet=ws.title, row=r,
                                 cols={k: f"{get_column_letter(c)}{r}" for k, c in cols.items() if c}))
    return Boq(path, items)
